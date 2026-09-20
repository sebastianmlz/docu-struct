"""Script para generar el documento Word formal (DOCX) explicativo de la arquitectura de DocuStruct."""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, hex_color):
    """Asigna color de fondo a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Asigna márgenes internos (padding) a una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="CONCEPTO CLAVE"):
    """Agrega una caja de llamada destacada."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Borde izquierdo azul grueso
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"[{title}] ")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(37, 99, 235)
    
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_docx():
    doc = Document()
    
    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # -------------------------------------------------------------
    # PORTADA / TÍTULO PRINCIPAL
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(6)
    t_run = title_p.add_run("DocuStruct")
    t_run.font.name = "Arial"
    t_run.font.size = Pt(28)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    sub_run = sub_p.add_run("Manual Técnico de Arquitectura, Ingeniería de Software y Flujo Operacional")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    # Metadatos del documento
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(24)
    m_run = meta_p.add_run("Versión: 0.1.0 (Producción Empresarial) | Cobertura de Pruebas: 93% | Licencia: Privada")
    m_run.font.name = "Arial"
    m_run.font.size = Pt(9.5)
    m_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("1. Introducción y Propósito del Sistema", level=1)
    
    p = doc.add_paragraph(
        "DocuStruct es un microservicio industrial de alta precisión diseñado para resolver un problema crítico "
        "en la gestión documental moderna: la extracción automatizada y estructuración de información contenida "
        "en archivos PDF complejos y heterogéneos (entre 1 y 10 páginas)."
    )
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph(
        "En entornos empresariales, gubernamentales y notariales, los documentos no son uniformes. Frecuentemente contienen "
        "tablas extensas que se cortan a través de múltiples páginas, credenciales de identidad insertadas en el texto "
        "(como RUT o DNI), timbres oficiales, firmas manuscritas y datos confidenciales tachados con marcador negro o cinta correctora. "
        "Las herramientas tradicionales de OCR fallan ante estos escenarios porque leen caracteres sueltos sin comprender la semántica visual. "
        "DocuStruct combina visión artificial avanzada (LLMs multimodales) con algoritmos deterministas en memoria para transformar "
        "esos documentos caóticos en estructuras JSON limpias, estandarizadas y listas para bases de datos."
    )
    p.paragraph_format.space_after = Pt(12)

    add_callout(
        doc,
        "Zero Disk I/O: El microservicio procesa los archivos PDF íntegramente en memoria RAM. "
        "Nunca se guarda una sola página o imagen en el disco del servidor, protegiendo la privacidad "
        "de datos sensibles y anulando vulnerabilidades de inyección en archivos temporales.",
        "PRINCIPIO ARQUITECTÓNICO"
    )

    doc.add_heading("2. Glosario Técnico Esencial: ¿Qué Significa Cada Cosa?", level=1)
    
    p = doc.add_paragraph(
        "Para comprender el funcionamiento interno de DocuStruct, a continuación se desglosan los conceptos "
        "tecnológicos clave que componen su arquitectura:"
    )
    p.paragraph_format.space_after = Pt(8)

    concepts = [
        ("FastAPI y ASGI", 
         "FastAPI es un framework web de alto rendimiento para Python. 'ASGI' (Asynchronous Server Gateway Interface) "
         "significa que el servidor puede atender múltiples peticiones concurrentes sin bloquearse mientras espera respuestas externas (como la API de IA)."),
        
        ("Rasterizado y DPI (pypdfium2)", 
         "Rasterizar es el proceso de convertir las instrucciones vectoriales de un PDF en una imagen compuesta por píxeles. "
         "'pypdfium2' es el motor ultrarrápido en C++ (el mismo de Google Chrome) que usamos en memoria. "
         "'DPI' (Dots Per Inch o Puntos Por Pulgada) mide la nitidez: 150 DPI fue elegido porque ofrece nitidez quirúrgica para leer números pequeños sin saturar los tokens de la IA."),
        
        ("Modelo Multimodal (Vision LLM)", 
         "Es un modelo de Inteligencia Artificial que no solo entiende texto, sino que tiene 'ojos'. Al recibir la imagen renderizada de la página, "
         "interpreta el orden de lectura, detecta firmas, analiza bordes de tablas y reconoce zonas censuradas."),
        
        ("LangChain LCEL", 
         "LangChain Expression Language. Es una sintaxis declarativa que permite encadenar componentes de software como tuberías (pipes `|`). "
         "Garantiza que la entrada se valide, el prompt se ensamble y la salida del modelo se fuerce a un esquema estricto de datos."),
        
        ("Pydantic v2", 
         "Es la librería estándar de validación de datos en Python. Si el modelo de IA intentara devolver un texto malformado, "
         "Pydantic lo rechaza de inmediato garantizando que el JSON resultante siempre respete los tipos definidos (enteros, listas, strings)."),
        
        ("Reconciliación Determinista", 
         "Significa que la unión de tablas divididas entre páginas no se deja al azar o a la 'imaginación' de la IA. "
         "Un algoritmo matemático en código Python puro compara encabezados, número de columnas y continuidad de filas para fusionarlas con exactitud matemática."),
        
        ("Idempotencia por Hash SHA-256", 
         "Si un usuario sube el mismo PDF dos veces, el sistema calcula su 'huella digital' criptográfica (SHA-256). "
         "Si la huella ya existe en la memoria caché, devuelve el resultado anterior en 5 milisegundos sin volver a consultar a la IA, ahorrando dinero y tiempo."),
        
        ("Control Anti-OOM y Concurrency Guard", 
         "OOM significa 'Out Of Memory' (servidor colgado por falta de RAM). El Concurrency Guard utiliza un semáforo asíncrono para permitir "
         "solo un número fijo de documentos pesados a la vez (por defecto 3). Si llega un 4to documento, se le responde de inmediato con HTTP 503 "
         "y una cabecera 'Retry-After: 5' (contrapresión), protegiendo la estabilidad del VPS."),
        
        ("Sliding Window Rate Limiter", 
         "Limitador de peticiones por ventana deslizante. Registra en memoria las marcas de tiempo de las peticiones de cada dirección IP. "
         "Si una IP realiza más de 10 peticiones en un minuto, es bloqueada con HTTP 429 para impedir ataques de denegación de servicio (DoS)."),
        
        ("Firma Binaria (Magic Bytes)", 
         "Los archivos informáticos tienen firmas ocultas en sus primeros bytes. Un PDF legítimo siempre inicia con `%PDF-`. "
         "DocuStruct inspecciona estos bytes antes de tocar cualquier otra cosa, bloqueando scripts maliciosos o virus camuflados con extensión `.pdf`."),
        
        ("Métricas Prometheus (/metrics)", 
         "Estándar de la industria para monitoreo de servidores. DocuStruct exporta métricas numéricas (peticiones totales, aciertos de caché, latencia) "
         "en texto plano para que herramientas como Grafana o Datadog vigilen la salud del servidor sin requerir programas pesados adicionales."),
        
        ("Trazabilidad con X-Request-ID", 
         "Cada solicitud recibe un código único (ej: `a1b2c3d4-...`). Ese código viaja en cada línea de log, permitiendo rastrear exactamente "
         "qué ocurrió con un documento específico entre cientos de peticiones simultáneas."),
        
        ("Docker y Usuario Sin Privilegios (Rootless)", 
         "El sistema se empaqueta en un contenedor aislado. Para máxima seguridad, no corre como administrador (root), "
         "sino bajo un usuario restringido llamado `appuser` (UID 10001) que no tiene permisos para alterar el sistema operativo anfitrión.")
    ]

    for title, desc in concepts:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_t = p_c.add_run(f"• {title}: ")
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        
        r_d = p_c.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_heading("3. Estructura de Carpetas y Archivos Explicada", level=1)
    
    p = doc.add_paragraph(
        "A continuación se detalla el rol exacto de cada directorio y archivo en el repositorio:"
    )
    p.paragraph_format.space_after = Pt(8)

    # Tabla de estructura
    table_files = [
        ("src/main.py", "Punto de entrada de la aplicación. Configura middlewares de seguridad, CORS, ciclo de vida (lifespan) y monta rutas web y API."),
        ("src/core/config.py", "Centraliza la configuración mediante Pydantic Settings. Lee variables de entorno (.env) con validación de tipos estrictos."),
        ("src/core/security.py", "Capa de defensa perimetral: validación de magic bytes (%PDF-), sanitizador de nombres de archivo agnóstico de SO, limitador de tasa y guardián de concurrencia."),
        ("src/core/cache.py", "Almacén en memoria RAM con algoritmo LRU (Least Recently Used) y TTL, indexado por hash criptográfico SHA-256 para evitar doble consumo de cuota."),
        ("src/core/telemetry.py", "Motor de métricas nativo que expone el formato estándar Prometheus en /metrics y JSON en /api/v1/telemetry."),
        ("src/core/logging.py", "Configuración de observabilidad estructurada con inyección de X-Request-ID vía ContextVars y formateador JSON para producción."),
        ("src/schemas/document.py", "Esquemas Pydantic v2 que definen las entidades de negocio: tablas, credenciales DNI/RUT, bloques de lectura y centinela de censura."),
        ("src/schemas/api_response.py", "Envoltura estándar para todas las respuestas HTTP (código de estado, mensaje, payload unificado y métricas de procesamiento)."),
        ("src/services/pdf_converter.py", "Servicio que toma los bytes del PDF y los convierte en imágenes PNG a 150 DPI en memoria mediante pypdfium2 con blindaje contra bombas de píxeles."),
        ("src/services/vision_extractor.py", "Orquestador de inferencia multimodal con LangChain LCEL. Administra prompts especializados, defensas contra inyección y reintentos con jitter aleatorio."),
        ("src/services/postprocessor.py", "Motor de reconciliación determinista. Une tablas partidas a través de páginas, detecta continuaciones implícitas y deduplica firmas institucionales."),
        ("src/api/v1/documents.py", "Controlador del endpoint POST /api/v1/documents/process. Orquesta la seguridad, caché, rasterizado y respuesta."),
        ("src/web/templates/index.html", "Interfaz de usuario forense con diseño Slate/Zinc sobrio, iconos vectoriales SVG limpios y sin emojis."),
        ("src/web/static/js/app.js", "Cliente web nativo en Vanilla JavaScript (ES6) para subida de archivos vía Fetch API y renderizado interactivo de resultados."),
        (".github/workflows/ci.yml", "Tubería de Integración Continua: linteo con Ruff, análisis de seguridad AST con Bandit, pruebas en Python 3.10/3.11 y smoke test de Docker."),
        (".github/workflows/cd.yml", "Tubería de Entrega Continua: compilación de imagen, escaneo de vulnerabilidades con Trivy y publicación en GitHub Container Registry (GHCR)."),
        ("Dockerfile", "Archivo de construcción de contenedor en múltiples etapas sobre Linux Debian Bookworm, ejecutado bajo el usuario appuser (10001)."),
        ("docker-compose.yml", "Orquestador de ejecución local y VPS con límites estrictos de CPU (2 núcleos) y memoria RAM (1 GB).")
    ]

    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    hdr.cells[0].text = "Archivo / Ruta"
    hdr.cells[1].text = "Función en el Sistema"
    set_cell_background(hdr.cells[0], "0F172A")
    set_cell_background(hdr.cells[1], "0F172A")
    hdr.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    hdr.cells[0].paragraphs[0].runs[0].bold = True
    hdr.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    hdr.cells[1].paragraphs[0].runs[0].bold = True

    for fpath, fdesc in table_files:
        row = tbl.add_row()
        row.cells[0].text = fpath
        row.cells[1].text = fdesc
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_margins(row.cells[0], top=80, bottom=80, left=120, right=120)
        set_cell_margins(row.cells[1], top=80, bottom=80, left=120, right=120)
        row.cells[0].paragraphs[0].runs[0].font.name = "Consolas"
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].paragraphs[0].runs[0].font.name = "Arial"
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_heading("4. Flujo Operacional Paso a Paso: De Principio a Fin", level=1)
    
    p = doc.add_paragraph(
        "Cuando un cliente (sea la interfaz web, una aplicación móvil o un script de automatización) envía un PDF a DocuStruct, "
        "la solicitud atraviesa una secuencia rigurosa de 11 pasos defensivos y de procesamiento:"
    )
    p.paragraph_format.space_after = Pt(8)

    steps = [
        ("Paso 1: Recepción e Inyección de Request-ID", 
         "La petición llega al servidor FastAPI. El middleware de observabilidad lee la cabecera 'X-Request-ID' o genera un identificador UUIDv4 único. "
         "Este código se fija en la variable de contexto asíncrono para que cualquier log emitido a partir de este instante lo contenga."),
        
        ("Paso 2: Cabeceras de Seguridad y Aislamiento de Navegador", 
         "El middleware de seguridad inyecta cabeceras defensivas: Content-Security-Policy estricta (bloquea recursos externos no autorizados), "
         "X-Frame-Options: DENY (evita ataques de Clickjacking), X-Content-Type-Options: nosniff y políticas de aislamiento de orígenes."),
        
        ("Paso 3: Control de Tasa (Rate Limiting en RAM)", 
         "Se identifica la dirección IP del cliente (soportando proxies reversos Nginx mediante X-Forwarded-For). "
         "Si la IP realizó más de 10 peticiones en el último minuto, se aborta con HTTP 429 Too Many Requests y se informa en cuántos segundos reintentar."),
        
        ("Paso 4: Sanitización Forense del Nombre de Archivo", 
         "Para neutralizar ataques de Directory Traversal (secuencias '../' o '..\\'), el nombre del archivo se normaliza de manera agnóstica "
         "entre Windows y Linux, extrayendo únicamente el nombre base y descartando caracteres de control o inyecciones."),
        
        ("Paso 5: Validación Temprana de Firma Binaria (Magic Bytes)", 
         "Se leen los primeros bytes del archivo en memoria. Si no comienzan con la firma ISO 32000-1 ('%PDF-'), la solicitud es rechazada de inmediato "
         "con HTTP 400 Bad Request. Esto detiene ejecutables (.exe), páginas HTML o scripts maliciosos antes de que lleguen al motor de PDFium."),
        
        ("Paso 6: Búsqueda en Caché Idempotente SHA-256 (Ahorro de Costos)", 
         "Se calcula el hash SHA-256 del binario y se combina con el modelo activo y el DPI. Si este documento exacto ya fue procesado y no ha expirado su TTL, "
         "se devuelve la respuesta previa en menos de 50 milisegundos con la cabecera 'X-Cache-Status: HIT'. Cero costo de IA y cero uso de semáforo."),
        
        ("Paso 7: Adquisición de Ranura en ConcurrencyGuard (Blindaje Anti-OOM)", 
         "Si no estaba en caché, el proceso solicita una ranura en el semáforo asíncrono de concurrencia. Si las 3 ranuras disponibles en memoria están ocupadas, "
         "se rechaza con HTTP 503 y 'Retry-After: 5', evitando que la memoria RAM del VPS se agote."),
        
        ("Paso 8: Rasterizado a 150 DPI en Memoria (Zero Disk I/O)", 
         "El servicio PDFConverterService invoca a pypdfium2. Cada página se convierte a PNG en un buffer de memoria RAM y se codifica en Base64. "
         "Se comprueba que las dimensiones no excedan 5000 puntos para neutralizar bombas de descompresión (pixel bombs)."),
        
        ("Paso 9: Inferencia Multimodal Secuencial con Contexto Inter-Página", 
         "Las páginas se envían secuencialmente al modelo multimodal (Gemini o GPT). Si la página 1 dejó una tabla abierta que continúa en la página 2, "
         "los encabezados y metadatos de esa tabla se inyectan automáticamente en el prompt de la página 2. Si la API externa arroja un error 429 transitorio, "
         "se activa el reintento automático con retroceso exponencial y jitter aleatorio."),
        
        ("Paso 10: Reconciliación Determinista en Postprocesador", 
         "El DocumentPostprocessor une los fragmentos de tablas divididas comparando encabezados y columnas, compensa celdas vacías para evitar desplazamiento, "
         "deduplica firmas repetidas y estructura todas las credenciales de identidad respetando el 'reading_order_index' original."),
        
        ("Paso 11: Indexación en Caché, Telemetría y Respuesta Exitosa", 
         "El resultado final se almacena en la caché LRU para futuras peticiones, se incrementan los contadores de telemetría Prometheus (/metrics), "
         "se libera la ranura del semáforo y se retorna el JSON unificado al cliente con código HTTP 200 OK y 'X-Cache-Status: MISS'.")
    ]

    for stitle, sdesc in steps:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(6)
        p_s.paragraph_format.space_after = Pt(2)
        r_st = p_s.add_run(stitle)
        r_st.bold = True
        r_st.font.name = "Arial"
        r_st.font.size = Pt(10.5)
        r_st.font.color.rgb = RGBColor(30, 64, 175)
        
        p_sd = doc.add_paragraph()
        p_sd.paragraph_format.space_after = Pt(6)
        r_sd = p_sd.add_run(sdesc)
        r_sd.font.name = "Arial"
        r_sd.font.size = Pt(9.5)
        r_sd.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_heading("5. Resolución de los Tres Casos Complejos del Negocio", level=1)

    p = doc.add_paragraph(
        "DocuStruct fue concebido específicamente para resolver con solidez tres escenarios donde los OCRs habituales fallan:"
    )
    p.paragraph_format.space_after = Pt(8)

    add_callout(
        doc,
        "Caso 1: Tablas Divididas entre Páginas (Continuous Tables)\n"
        "Cuando una tabla de licitación o presupuesto supera una página, la página siguiente suele omitir los encabezados "
        "o arrancar directamente en una fila de datos. DocuStruct inyecta el esquema de la tabla de la página anterior como contexto, "
        "identifica la continuidad y las fusiona en una única matriz estructurada sin duplicar títulos ni perder filas.",
        "CASO DE NEGOCIO 1"
    )

    add_callout(
        doc,
        "Caso 2: Documentos Polimórficos y Preservación del Orden de Lectura\n"
        "En un mismo folio pueden coexistir: un membrete legal, un bloque de credencial de identidad (Nombre, RUN/DNI), "
        "una tabla financiera y finalmente firmas con timbres notariales. DocuStruct asigna un 'reading_order_index' secuencial "
        "a cada bloque, permitiendo reconstruir la secuencia lógica exacta del documento sin mezclar textos.",
        "CASO DE NEGOCIO 2"
    )

    add_callout(
        doc,
        "Caso 3: Detección Forense de Datos Censurados (Redaction Handling)\n"
        "Si un documento oficial tiene nombres, montos o RUTs tachados con marcador negro o corrector líquido, "
        "el sistema no desplaza las columnas hacia la izquierda (lo que provocaría que los montos queden en la columna de nombres). "
        "El modelo detecta la tachadura visual y coloca exactamente en esa celda el centinela '[DATO_CENSURADO]', preservando la integridad tabular.",
        "CASO DE NEGOCIO 3"
    )

    doc.add_heading("6. Guía de Operación y Monitoreo en Producción", level=1)
    
    p = doc.add_paragraph(
        "Para operar el microservicio en un entorno local o servidor VPS dedicado, se disponen de los siguientes comandos y endpoints:"
    )
    p.paragraph_format.space_after = Pt(6)

    ops = [
        ("Ejecución con Docker Compose:", "docker compose up -d --build"),
        ("Verificación de Salud:", "curl -s http://127.0.0.1:8000/api/v1/health"),
        ("Métricas de Prometheus:", "curl -s http://127.0.0.1:8000/metrics"),
        ("Telemetría en JSON:", "curl -s http://127.0.0.1:8000/api/v1/telemetry"),
        ("Ejecutar Suite de Tests (81 pruebas):", "pytest -v --cov=src --cov-report=term-missing"),
        ("Auditoría Estática de Seguridad:", "bandit -r src/ -ll"),
        ("Linteo de Código:", "ruff check src/ tests/")
    ]

    tbl_op = doc.add_table(rows=1, cols=2)
    tbl_op.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_op = tbl_op.rows[0]
    h_op.cells[0].text = "Acción / Verificación"
    h_op.cells[1].text = "Comando de Terminal"
    set_cell_background(h_op.cells[0], "1E293B")
    set_cell_background(h_op.cells[1], "1E293B")
    h_op.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h_op.cells[0].paragraphs[0].runs[0].bold = True
    h_op.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h_op.cells[1].paragraphs[0].runs[0].bold = True

    for act, cmd in ops:
        row = tbl_op.add_row()
        row.cells[0].text = act
        row.cells[1].text = cmd
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_margins(row.cells[0], top=60, bottom=60, left=100, right=100)
        set_cell_margins(row.cells[1], top=60, bottom=60, left=100, right=100)
        row.cells[0].paragraphs[0].runs[0].font.name = "Arial"
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        row.cells[1].paragraphs[0].runs[0].font.name = "Consolas"
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Conclusión final
    p_fin = doc.add_paragraph(
        "Conclusión: DocuStruct no es un prototipo experimental, sino un microservicio diseñado bajo los más altos estándares "
        "de ingeniería de software: tipado estricto, 93% de cobertura de pruebas, cero vulnerabilidades estáticas, resiliencia ante "
        "caídas de cuota, blindaje anti-OOM para servidores VPS y observabilidad integral para entornos corporativos."
    )
    p_fin.paragraph_format.space_before = Pt(12)
    p_fin.runs[0].font.name = "Arial"
    p_fin.runs[0].font.size = Pt(9.5)
    p_fin.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    out_path = os.path.abspath("DocuStruct_Manual_Tecnico_Arquitectura.docx")
    doc.save(out_path)
    print(f"Documento Word generado exitosamente en: {out_path}")

if __name__ == "__main__":
    build_docx()
